from __future__ import annotations
import csv
from typing import List

from kivy.uix.boxlayout import BoxLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.gridlayout import GridLayout
from kivy.uix.label import Label
from kivy.metrics import dp
from kivy.properties import StringProperty, NumericProperty, ObjectProperty
from kivy.clock import Clock

from ui.styles import colors

class SpreadsheetPreview(BoxLayout):
    """Einfache CSV/XLSX Vorschau. Liest aus Datei und zeigt die ersten N Zeilen."""
    source = StringProperty("")         # Pfad zu .csv oder .xlsx
    filetype = StringProperty("csv")    # "csv" | "xlsx"
    delimiter = StringProperty(";")
    max_rows = NumericProperty(200)
    max_cols = NumericProperty(50)

    _sv = ObjectProperty(None, allownone=True)
    _grid = ObjectProperty(None, allownone=True)

    def __init__(self, **kwargs):
        super().__init__(orientation="vertical", **kwargs)
        sv = ScrollView(do_scroll_x=True, do_scroll_y=True, bar_width=dp(6))
        grid = GridLayout(cols=1, spacing=dp(2), size_hint=(None, None))
        grid.bind(minimum_height=grid.setter("height"))
        grid.bind(minimum_width=grid.setter("width"))
        sv.add_widget(grid)
        self._sv = sv
        self._grid = grid
        self.add_widget(sv)
        Clock.schedule_once(lambda dt: self._reload(), 0)

    def on_source(self, *_):
        if not self._grid:
            Clock.schedule_once(lambda dt: self._reload(), 0)
            return
        self._reload()

    def on_filetype(self, *_):
        if not self._grid:
            Clock.schedule_once(lambda dt: self._reload(), 0)
            return
        self._reload()

    def _reload(self):
        if not self.source or not self._grid or not self._sv:
            return
        rows: List[List[str]] = []
        try:
            if self.filetype == "csv":
                with open(self.source, "r", encoding="utf-8", newline="") as f:
                    r = csv.reader(f, delimiter=self.delimiter)
                    for i, row in enumerate(r):
                        if i >= self.max_rows:
                            break
                        rows.append([str(x) for x in row[: self.max_cols]])
            else:
                from openpyxl import load_workbook
                wb = load_workbook(self.source, read_only=True, data_only=True)
                ws = wb.active
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= self.max_rows:
                        break
                    rows.append([("" if c is None else str(c)) for c in row[: self.max_cols]])
        except Exception as e:
            self._grid.clear_widgets()
            self._grid.cols = 1
            self._grid.add_widget(Label(text=f"Vorschau konnte nicht geladen werden: {e}",
                                        color=colors.TEXT_COLOR))
            self._grid.width = self._sv.width
            return

        self._build_table(rows)

    def _build_table(self, rows: List[List[str]]):
        if not self._grid or not self._sv:
            return
        self._grid.clear_widgets()
        if not rows:
            self._grid.cols = 1
            self._grid.add_widget(Label(text="Keine Daten", color=colors.TEXT_COLOR))
            self._grid.width = self._sv.width
            return

        cols = max(1, max(len(r) for r in rows))
        self._grid.cols = cols

        # Zellenbreite grob schätzen
        base_w = dp(140)
        self._grid.width = base_w * cols

        # Kopfzeile
        for j, cell in enumerate(rows[0]):
            self._grid.add_widget(Label(
                text=str(cell),
                bold=True,
                color=colors.TEXT_COLOR,
                size_hint=(None, None),
                size=(base_w, dp(28)),
                halign="left", valign="middle",
                text_size=(base_w - dp(10), dp(28))
            ))

        # Rest
        for r in rows[1:]:
            for j in range(cols):
                txt = r[j] if j < len(r) else ""
                self._grid.add_widget(Label(
                    text=str(txt),
                    color=colors.TEXT_COLOR,
                    size_hint=(None, None),
                    size=(base_w, dp(26)),
                    halign="left", valign="middle",
                    text_size=(base_w - dp(10), dp(26))
                ))
