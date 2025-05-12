from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import os
from tkinter import filedialog

def export_to_excel(devices, networks):
    """Exportiert die Scan-Ergebnisse als Excel-Datei"""
    try:
        # Erstelle neue Excel-Arbeitsmappe
        wb = Workbook()
        
        # Erstelle Geräte-Worksheet
        ws_devices = wb.active
        ws_devices.title = "Geräte"
        
        # Überschriften für Geräte
        headers_devices = ["Device ID", "Vendor ID", "IP-Adresse", "Port", "Netzwerk"]
        ws_devices.append(headers_devices)
        
        # Style für Überschriften
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="00CCCC", end_color="00CCCC", fill_type="solid")
        
        # Überschriften formatieren
        for cell in ws_devices[1]:
            cell.font = header_font
            cell.fill = header_fill
        
        # Geräte-Daten einfügen
        for device_id, device_info in devices.items():
            row = [
                device_id,
                device_info.get('vendor_id', ''),
                device_info.get('address', ''),
                device_info.get('port', ''),
                device_info.get('network', '')
            ]
            ws_devices.append(row)
        
        # Erstelle Netzwerk-Worksheet
        ws_networks = wb.create_sheet(title="Netzwerke")
        
        # Überschriften für Netzwerke
        headers_networks = ["Netzwerk ID", "Beschreibung"]
        ws_networks.append(headers_networks)
        
        # Überschriften formatieren
        for cell in ws_networks[1]:
            cell.font = header_font
            cell.fill = header_fill
        
        # Netzwerk-Daten einfügen
        for network_id, network_info in networks.items():
            row = [network_id, network_info.get('description', '')]
            ws_networks.append(row)
        
        # Spaltenbreiten automatisch anpassen
        for ws in [ws_devices, ws_networks]:
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Datei speichern
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel-Dateien", "*.xlsx")],
            title="Excel-Datei speichern"
        )
        
        if file_path:
            wb.save(file_path)
            return True, f"Daten wurden erfolgreich nach {file_path} exportiert"
        else:
            return False, "Export abgebrochen"
            
    except Exception as e:
        return False, f"Fehler beim Export: {str(e)}"